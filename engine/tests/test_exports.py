import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from next_signal.export.csvs import build_zip, table_rows
from next_signal.export.xlsx import build_workbook
from next_signal.schemas import Report

FIXTURE = Path(__file__).parent / "fixtures" / "report_min.json"


@pytest.fixture
def report() -> Report:
    return Report.model_validate(json.loads(FIXTURE.read_text(encoding="utf-8")))


def test_table_rows_evidence_has_header_and_data(report):
    rows = table_rows(report, "evidence")
    assert rows[0][0] == "ev_id"
    assert len(rows) == 1 + len(report.discovery.evidence)


def test_table_rows_unknown_table_raises(report):
    with pytest.raises(ValueError):
        table_rows(report, "not_a_real_table")


def test_build_zip_contains_ten_csvs(report, tmp_path):
    out = build_zip(report, tmp_path / "export.zip")
    assert out.exists()
    with zipfile.ZipFile(out) as zf:
        names = zf.namelist()
        assert len(names) == 10
        assert "evidence.csv" in names
        content = zf.read("evidence.csv")
        assert content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM present


def test_csv_content_has_no_dashes_in_quotes(report):
    rows = table_rows(report, "evidence")
    for row in rows[1:]:
        quote = row[-1]
        assert "—" not in quote and "–" not in quote


def test_build_workbook_has_discovery_plus_ten_sheets(report, tmp_path):
    out = build_workbook(report, tmp_path / "export.xlsx")
    assert out.exists()
    wb = openpyxl.load_workbook(out)
    assert "Discovery" in wb.sheetnames
    assert len(wb.sheetnames) == 11  # Discovery + 10 tables
    assert wb["Discovery"]["A2"].value == report.discovery.sentence


def test_workbook_header_row_styled(report, tmp_path):
    out = build_workbook(report, tmp_path / "export.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["Evidence"]
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold is True
