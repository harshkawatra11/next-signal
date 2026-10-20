"""Excel workbook export: one sheet per table, header row bold white on signal red, frozen header,
autofilter, hyperlinks on URL columns.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..schemas import Report
from .csvs import _TABLES, table_rows

HEADER_FILL = PatternFill(start_color="C8102E", end_color="C8102E", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet_title(name: str) -> str:
    titles = {
        "evidence": "Evidence", "numbers": "Numbers", "timeline": "Timeline", "hypotheses": "Hypotheses",
        "tests": "Tests", "ach": "ACH", "contradictions": "Contradictions", "sources": "Sources",
        "metrics": "Metrics", "anomalies": "Anomalies",
    }
    return titles.get(name, name.title())[:31]


def build_workbook(report: Report, out_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    disc_ws = wb.create_sheet("Discovery")
    d = report.discovery
    disc_ws["A1"] = "The Discovery"
    disc_ws["A1"].font = Font(bold=True, size=14)
    disc_ws["A2"] = d.sentence if d else "No discovery survived the challenge stage."
    disc_ws["A2"].alignment = Alignment(wrap_text=True)
    disc_ws.column_dimensions["A"].width = 100
    if d:
        disc_ws["A4"] = "Strength"
        disc_ws["A4"].font = Font(bold=True)
        disc_ws["A5"] = f"{d.strength.label} ({d.strength.total}/10)"
        disc_ws["A6"] = "; ".join(d.strength.reasoning)
        disc_ws["A6"].alignment = Alignment(wrap_text=True)
        disc_ws["A8"] = "Why it matters"
        disc_ws["A8"].font = Font(bold=True)
        for i, bullet in enumerate(d.why_it_matters):
            disc_ws[f"A{9 + i}"] = f"- {bullet}"
        disc_ws["A14"] = "Counterargument"
        disc_ws["A14"].font = Font(bold=True)
        disc_ws["A15"] = d.counterargument
        disc_ws["A17"] = "Falsification condition"
        disc_ws["A17"].font = Font(bold=True)
        disc_ws["A18"] = d.falsification.statement

    for table in _TABLES:
        rows = table_rows(report, table)
        ws = wb.create_sheet(_sheet_title(table))
        for row in rows:
            ws.append(row)
        if rows:
            for col_idx, header in enumerate(rows[0], start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col_idx, header in enumerate(rows[0], start=1):
                max_len = max((len(str(r[col_idx - 1])) for r in rows if col_idx - 1 < len(r)), default=10)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(80, max(10, max_len + 2))
            if "url" in [h.lower() for h in rows[0]]:
                url_col = [h.lower() for h in rows[0]].index("url") + 1
                for r in range(2, len(rows) + 1):
                    cell = ws.cell(row=r, column=url_col)
                    if cell.value:
                        cell.hyperlink = cell.value
                        cell.font = Font(color="0563C1", underline="single")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
